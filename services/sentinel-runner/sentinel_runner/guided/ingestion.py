"""Bounded parsing without execution; source rows and passages are never synthesized."""
from __future__ import annotations

import hashlib
import io
import json
import math
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

MAX_BYTES = 2_000_000
MAX_EXPANDED_BYTES = 16_000_000
MAX_ROWS = 5_000
MAX_COLUMNS = 64
MAX_PASSAGES = 256
PARSER_VERSION = "eidos.data.v1.0"
LABEL = re.compile(r"(^|[_\s-])(label|attack|class|ground.?truth|target.?label)([_\s-]|$)", re.I)
TIME = re.compile(r"timestamp|datetime|event.?time|^date$|^time$", re.I)
ENTITY = re.compile(r"(^|_)(host|service|customer|entity|session|device)(_id)?$", re.I)


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def plain(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(k): plain(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, np.ndarray)):
        return [plain(v) for v in value]
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, (float, np.floating)):
        return float(value) if math.isfinite(value) else None
    if isinstance(value, (int, np.integer)):
        return int(value)
    if isinstance(value, (pd.Timestamp,)):
        return value.isoformat()
    if isinstance(value, (str, bool)):
        return value
    return str(value)


class VisibleHTML(HTMLParser):
    """Scripts/styles/navigation are data we deliberately exclude, never execute."""
    excluded = {"script", "style", "nav", "header", "footer", "noscript", "svg", "form"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.stack: list[str] = []
        self.parts: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag in self.excluded:
            self.stack.append(tag)
        elif tag in {"p", "div", "br", "li", "h1", "h2", "h3", "article"} and not self.stack:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.stack:
            # Malformed markup never makes hidden script text visible early.
            self.stack = self.stack[:self.stack.index(tag)]
        elif tag in {"p", "div", "li", "article"} and not self.stack:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.stack:
            self.parts.append(data)


def passages(pages: list[tuple[int, str]], source: str) -> tuple[list[dict], list[str]]:
    result, seen, warnings = [], set(), []
    for page, text in pages:
        # Short, non-overlapping chunks avoid encoder truncation; all paragraphs count.
        for paragraph_no, paragraph in enumerate(re.split(r"\n\s*\n|\n", text)):
            paragraph = re.sub(r"\s+", " ", paragraph).strip()
            if not paragraph:
                continue
            words = paragraph.split()
            for start in range(0, len(words), 100):
                chunk = " ".join(words[start:start + 100])
                content_hash = digest(chunk.encode())
                if content_hash in seen:
                    warnings.append(f"Duplicate passage excluded at page {page}, paragraph {paragraph_no + 1}.")
                    continue
                seen.add(content_hash)
                if len(result) >= MAX_PASSAGES:
                    raise ValueError("PASSAGE_LIMIT: maximum 256 passages; split this document.")
                result.append({"id": f"p{page}-{paragraph_no + 1}-{start}", "source": source,
                               "page": page, "paragraph": paragraph_no + 1, "wordOffset": start,
                               "text": chunk, "sha256": content_hash})
    if not result:
        raise ValueError("NO_EXTRACTABLE_TEXT: image-only PDFs require OCR and are not supported.")
    return result, warnings


def log_rows(text: str) -> list[dict]:
    rows = []
    for line_number, line in enumerate(text.splitlines(), 1):
        if not line.strip():
            continue
        timestamp = re.search(r"\d{4}-\d{2}-\d{2}[T ][\d:.]+(?:Z|[+-]\d\d:\d\d)?", line)
        attributes = dict(re.findall(r"\b([A-Za-z_][\w.-]*)=([^\s]+)", line))
        template = re.sub(r"\b[0-9a-f]{8}-[0-9a-f-]{27,}\b", "<id>", line, flags=re.I)
        template = re.sub(r"\b(?:\d{1,3}\.){3}\d{1,3}\b", "<ip>", template)
        if timestamp:
            template = template.replace(timestamp.group(), "<timestamp>")
        template = re.sub(r"(?<![a-zA-Z])[-+]?\d+(?:\.\d+)?", "<number>", template)
        rows.append({"timestamp": timestamp.group() if timestamp else None, "template": template,
                     "message": line, "line": line_number, **attributes})
    return rows


def parse(data: bytes, filename: str, source: str | None = None) -> dict:
    if not data or len(data) > MAX_BYTES:
        raise ValueError("UPLOAD_SIZE: file must contain 1 to 2,000,000 bytes.")
    source = source or Path(filename).name
    extension = Path(filename).suffix.lower()
    warnings: list[str] = []
    pages = None
    if extension == ".csv":
        # Read strings first: IDs with leading zeros and original values survive.
        table = pd.read_csv(io.BytesIO(data), dtype=str, keep_default_na=False, nrows=MAX_ROWS + 1)
    elif extension == ".xlsx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(f.file_size for f in entries) > MAX_EXPANDED_BYTES:
                raise ValueError("EXPANDED_SIZE: workbook exceeds decompression limit.")
            if any("vbaProject" in f.filename for f in entries):
                raise ValueError("MACROS_UNSUPPORTED: supply a values-only workbook.")
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        if len(book.worksheets) != 1:
            raise ValueError("MULTIPLE_SHEETS: export the intended sheet separately.")
        sheet = book.worksheets[0]
        if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > MAX_COLUMNS:
            raise ValueError("TABLE_LIMIT: maximum 5,000 rows and 64 columns.")
        cells = list(sheet.iter_rows())
        if any(c.data_type == "f" for row in cells for c in row):
            raise ValueError("FORMULAS_UNSUPPORTED: export computed values before uploading.")
        table = pd.DataFrame([[c.value for c in row] for row in cells[1:]], columns=[c.value for c in cells[0]])
        book.close()
    elif extension == ".parquet":
        import pyarrow.parquet as pq
        parquet = pq.ParquetFile(io.BytesIO(data))
        if parquet.metadata.num_rows > MAX_ROWS or parquet.metadata.num_columns > MAX_COLUMNS or sum(
            parquet.metadata.row_group(i).total_byte_size for i in range(parquet.metadata.num_row_groups)
        ) > MAX_EXPANDED_BYTES:
            raise ValueError("TABLE_LIMIT: Parquet exceeds row/column/decompression limits.")
        table = parquet.read().to_pandas()
    elif extension in {".json", ".jsonl", ".ndjson"}:
        text = data.decode("utf-8-sig", errors="strict")
        records = json.loads(text) if extension == ".json" else [json.loads(line) for line in text.splitlines() if line.strip()]
        if not isinstance(records, list) or not records or not all(isinstance(row, dict) for row in records):
            raise ValueError("JSON_RECORDS_REQUIRED: expected an array or one object per line.")
        if len(records) > MAX_ROWS or any(isinstance(v, (dict, list)) for row in records for v in row.values()):
            raise ValueError("FLAT_RECORDS_REQUIRED: maximum 5,000 flat records.")
        table = pd.DataFrame(records)
    elif extension == ".log":
        table = pd.DataFrame(log_rows(data.decode("utf-8-sig", errors="strict")))
        warnings.append("Log templates replace timestamps, addresses and changing numeric IDs; originals remain in message.")
    elif extension == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted or len(reader.pages) > 50:
            raise ValueError("PDF_LIMIT: unencrypted text PDFs with at most 50 pages are supported.")
        pages = []
        for i, page in enumerate(reader.pages):
            contents = page.get_contents()
            if contents and len(contents.get_data()) > MAX_EXPANDED_BYTES:
                raise ValueError("EXPANDED_SIZE: PDF page exceeds extraction budget.")
            text = page.extract_text() or ""
            if not text.strip():
                warnings.append(f"Page {i + 1} has no extractable text; images were not interpreted.")
            pages.append((i + 1, text))
    elif extension in {".html", ".htm", ".txt", ".md"}:
        text = data.decode("utf-8-sig", errors="strict")
        if extension in {".html", ".htm"}:
            parser = VisibleHTML()
            parser.feed(text)
            text = "".join(parser.parts)
            warnings.append("Page scripts, styles, navigation and forms excluded; no page code executed.")
        pages = [(1, text)]
    else:
        raise ValueError("UNSUPPORTED_FORMAT: use CSV, XLSX, Parquet, JSON/JSONL, LOG, text PDF, HTML, TXT or MD.")
    base = {"schema": PARSER_VERSION, "source": source, "filename": Path(filename).name,
            "sha256": digest(data), "bytes": len(data), "warnings": warnings}
    if pages is not None:
        chunks, notes = passages(pages, source)
        return {**base, "kind": "documents", "passages": chunks, "rowCount": len(chunks),
                "warnings": warnings + notes, "chronology": "unordered", "columns": []}
    if not 1 <= len(table) <= MAX_ROWS or not 1 <= len(table.columns) <= MAX_COLUMNS:
        raise ValueError("TABLE_LIMIT: expected 1–5,000 rows and 1–64 columns.")
    names = [str(c).strip() for c in table.columns]
    if any(not c or len(c) > 100 for c in names) or len(set(names)) != len(names):
        raise ValueError("INVALID_COLUMNS: column names must be unique and 1–100 characters.")
    table.columns = names
    records = plain(table.to_dict(orient="records"))
    columns = []
    for name in names:
        values = table[name].replace(r"^\s*$", np.nan, regex=True)
        numeric = pd.to_numeric(values, errors="coerce")
        present = int(values.notna().sum())
        finite = np.isfinite(numeric)
        numeric_count = int(finite.sum())
        candidate_time = bool(TIME.search(name))
        columns.append({"name": name, "type": "timestamp_candidate" if candidate_time else "number" if present and numeric_count == present else "category",
                        "missing": len(values) - present, "invalidNumeric": int((values.notna() & ~finite).sum()),
                        "distinct": int(values.nunique()), "constant": values.nunique() <= 1,
                        "labelCandidate": bool(LABEL.search(name)), "entityCandidate": bool(ENTITY.search(name)),
                        "idCandidate": bool(re.search(r"(^id$|_id$|^line$|^index$)", name, re.I)),
                        "examples": plain(values.dropna().head(3).tolist())})
    categories = {c["name"]: [{"value": str(k), "count": int(v)} for k, v in table[c["name"]].value_counts(dropna=False).head(20).items()]
                  for c in columns if c["type"] == "category" and not c["labelCandidate"] and not c["idCandidate"]}
    if extension == ".log" and not any(c["type"] == "number" and not c["idCandidate"] and not c["constant"] for c in columns):
        chunks, notes = passages([(i + 1, row["template"]) for i, row in enumerate(records)], source)
        for chunk in chunks:
            chunk["originalRecordIds"] = [f"r{i + 1}" for i, row in enumerate(records) if chunk["text"] in row["template"]]
            chunk["originalExample"] = records[chunk["page"] - 1]["message"]
        return {**base, "kind": "documents", "representation": "log_templates", "passages": chunks,
                "records": records, "recordIds": [f"r{i + 1}" for i in range(len(records))], "columns": [],
                "rowCount": len(records), "chronology": "unordered", "categories": categories,
                "warnings": warnings + notes + ["No varying numeric measurement: semantic template analysis; no temporal forecast."]}
    return {**base, "kind": "table", "records": records, "columns": columns, "rowCount": len(records), "categories": categories,
            "duplicates": int(table.duplicated().sum()), "chronology": "unconfirmed",
            "recordIds": [f"r{i + 1}" for i in range(len(records))]}


def confirm(dataset: dict, mapping: dict) -> dict:
    """Validate meaning before analysis; do not derive chronology from file order."""
    if dataset["kind"] == "documents":
        return {"version": PARSER_VERSION, "mode": "unordered", "reference": "document collection", "units": {}, "features": []}
    by_name = {c["name"]: c for c in dataset["columns"]}
    features = mapping.get("features", [])
    labels = set(mapping.get("labels", [])) | {c["name"] for c in dataset["columns"] if c["labelCandidate"]}
    timestamp, entity = mapping.get("timestamp"), mapping.get("entity")
    session = mapping.get("session")
    exclusions = labels | {timestamp, entity, session}
    if not isinstance(features, list) or not features or len(features) > 16 or len(set(features)) != len(features):
        raise ValueError("FEATURES_REQUIRED: choose 1–16 distinct numeric measurements.")
    for feature in features:
        c = by_name.get(feature)
        if not c or feature in exclusions or c["idCandidate"] or c["constant"] or c["type"] != "number":
            raise ValueError(f"INVALID_FEATURE: {feature} is a label, ID, constant or nonnumeric measurement.")
    for column in [timestamp, entity, session, *labels]:
        if column and column not in by_name:
            raise ValueError(f"UNKNOWN_COLUMN: {column}")
    if not str(mapping.get("meaning", "")).strip() or not str(mapping.get("reference", "")).strip():
        raise ValueError("MEANING_REQUIRED: confirm what one row represents and the reference group.")
    if mapping.get("missing") not in {"exclude", "prefix_median"}:
        raise ValueError("MISSING_POLICY_REQUIRED: choose exclusion or calibration-prefix median imputation.")
    mode = "temporal" if timestamp else "unordered"
    dates, invalid, out_of_order, duplicate_times = [], 0, 0, 0
    previous, seen = {}, set()
    for row in dataset["records"]:
        if timestamp:
            try:
                t = parse_time(row.get(timestamp), mapping.get("timezone"))
                key = (str(row.get(entity)) if entity else "all", str(row.get(session)) if session else "all")
                if key in previous and t < previous[key]:
                    out_of_order += 1
                if (key, t) in seen:
                    duplicate_times += 1
                seen.add((key, t))
                previous[key] = t
                dates.append(t)
            except (ValueError, TypeError):
                invalid += 1
    if timestamp and (invalid or duplicate_times):
        raise ValueError(f"TIMESTAMP_QUALITY: {invalid} invalid and {duplicate_times} duplicate entity/session timestamps; repair or explicitly use unordered mode.")
    if out_of_order and mapping.get("ordering") != "sort":
        raise ValueError("ORDERING_CONFIRMATION: choose explicit per-entity chronological sorting.")
    units = mapping.get("units", {})
    if not isinstance(units, dict) or any(not isinstance(units.get(f), str) or not 1 <= len(units[f]) <= 40 for f in features):
        raise ValueError("UNITS_REQUIRED: specify units, or explicitly enter unknown, for each measurement.")
    return {**mapping, "features": features, "labels": sorted(labels), "version": PARSER_VERSION,
            "mode": mode, "coverage": {"from": min(dates) if dates else None, "to": max(dates) if dates else None},
            "outOfOrder": out_of_order, "invalidTimestamps": invalid, "duplicateTimestamps": duplicate_times,
            "excludedColumns": [c for c in by_name if c not in features], "units": units}


def parse_time(value, timezone: str | None) -> str:
    if not isinstance(value, str) or not re.match(r"^\d{4}-\d\d-\d\d[T ]\d\d:\d\d", value):
        raise ValueError("Use ISO-8601 timestamps; numeric epochs and ambiguous dates require conversion.")
    t = pd.Timestamp(value)
    if t.tzinfo is None:
        if not timezone:
            raise ValueError("Confirm an IANA time zone for timestamps without offsets.")
        t = t.tz_localize(timezone, ambiguous="raise", nonexistent="raise")
    return t.tz_convert("UTC").isoformat()
